"""
test_autoreport.py
==================
Suite de pruebas unitarias automatizadas con unittest para verificar la integridad,
robustez y precisión de cada componente de AutoReport_PyME:
1. Generación de datos sintéticos con anomalías.
2. Motor de limpieza, auditoría y cálculo de KPIs.
3. Exportadores ejecutivos a Excel (.xlsx) y PDF (.pdf) en memoria.
"""

from __future__ import annotations

import io
import os
import tempfile
import unittest

import numpy as np
import openpyxl
import pandas as pd

from cleaning_engine import clean_sales_data
from report_builder import export_to_excel_bytes, export_to_pdf_bytes
from sample_generator import generate_dirty_sales_data


class TestSampleGenerator(unittest.TestCase):
    """Pruebas unitarias para el generador de datos sucios de prueba."""

    def setUp(self):
        self.df_raw = generate_dirty_sales_data(filepath=None, n_rows=280, seed=123)

    def test_minimum_rows_and_columns(self):
        """Verifica que se generen suficientes registros y las columnas requeridas."""
        self.assertGreaterEqual(len(self.df_raw), 250)
        expected_cols = {
            "ID_Transaccion",
            "Fecha",
            "Sucursal",
            "Vendedor",
            "Categoria",
            "Producto",
            "Cantidad",
            "Precio_Unitario",
            "Costo_Unitario",
        }
        self.assertTrue(expected_cols.issubset(set(self.df_raw.columns)))

    def test_intentional_anomalies_presence(self):
        """Verifica que el dataset contenga anomalías deliberadas para auditoría."""
        # 1. Existencia de duplicados
        num_duplicates = self.df_raw.duplicated().sum()
        self.assertGreater(num_duplicates, 0, "Debe haber duplicados intencionales.")

        # 2. Existencia de valores nulos
        num_nulls = self.df_raw.isna().sum().sum()
        self.assertGreater(num_nulls, 0, "Debe haber celdas nulas intencionales.")

        # 3. Textos con espacios residuales o casing sucio
        sucursales_str = self.df_raw["Sucursal"].dropna().astype(str).tolist()
        has_spaces_or_lower = any(" " in s or s != s.strip() or s.islower() for s in sucursales_str)
        self.assertTrue(has_spaces_or_lower, "Debe haber textos desalineados en sucursales.")

    def test_excel_file_export_on_disk(self):
        """Verifica que la función guarde físicamente el archivo cuando se solicita."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            df = generate_dirty_sales_data(filepath=tmp_path, n_rows=50, seed=99)
            self.assertTrue(os.path.exists(tmp_path))
            self.assertGreater(os.path.getsize(tmp_path), 1000)
            loaded_df = pd.read_excel(tmp_path)
            self.assertEqual(len(loaded_df), len(df))
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)


class TestCleaningEngine(unittest.TestCase):
    """Pruebas unitarias para el motor de saneamiento y KPIs."""

    def setUp(self):
        # Generar un dataset sucio controlado
        self.df_raw = generate_dirty_sales_data(filepath=None, n_rows=300, seed=42)
        self.result = clean_sales_data(self.df_raw)
        self.df_clean = self.result["cleaned_df"]
        self.audit = self.result["audit_metrics"]
        self.kpis = self.result["financial_kpis"]

    def test_dictionary_structure(self):
        """Verifica que el resultado contenga todas las claves requeridas."""
        self.assertIn("cleaned_df", self.result)
        self.assertIn("audit_metrics", self.result)
        self.assertIn("financial_kpis", self.result)
        self.assertIsInstance(self.df_clean, pd.DataFrame)
        self.assertIsInstance(self.audit, dict)
        self.assertIsInstance(self.kpis, dict)

    def test_deduplication_and_nulls_removal(self):
        """Verifica la eliminación estricta de duplicados y nulos críticos."""
        # Cero duplicados en el dataset limpio
        self.assertEqual(self.df_clean.duplicated().sum(), 0)

        # Cero nulos en columnas críticas
        critical_cols = ["ID_Transaccion", "Fecha", "Sucursal", "Vendedor", "Cantidad", "Precio_Unitario", "Costo_Unitario"]
        for col in critical_cols:
            self.assertEqual(self.df_clean[col].isna().sum(), 0, f"Columna {col} no debe contener nulos.")

    def test_text_normalization(self):
        """Verifica la canonicidad de sucursales y vendedores."""
        valid_branches = {"Monterrey", "Guadalajara", "CDMX"}
        unique_branches = set(self.df_clean["Sucursal"].unique())
        self.assertTrue(
            unique_branches.issubset(valid_branches),
            f"Las sucursales deben ser exclusivamente {valid_branches}, encontradas: {unique_branches}",
        )

        # Verificar que ningún vendedor tenga espacios en blanco adelante o atrás
        for seller in self.df_clean["Vendedor"].unique():
            self.assertEqual(seller, seller.strip())

    def test_financial_calculations_accuracy(self):
        """Verifica la exactitud matemática de Venta_Total, Costo_Total, Utilidad y Margen_Pct."""
        calculated_sales = (self.df_clean["Cantidad"] * self.df_clean["Precio_Unitario"]).round(2)
        np.testing.assert_allclose(self.df_clean["Venta_Total"], calculated_sales, rtol=1e-5)

        calculated_cost = (self.df_clean["Cantidad"] * self.df_clean["Costo_Unitario"]).round(2)
        np.testing.assert_allclose(self.df_clean["Costo_Total"], calculated_cost, rtol=1e-5)

        calculated_profit = (self.df_clean["Venta_Total"] - self.df_clean["Costo_Total"]).round(2)
        np.testing.assert_allclose(self.df_clean["Utilidad"], calculated_profit, rtol=1e-5)

        # Margen %
        calculated_margin = ((self.df_clean["Utilidad"] / self.df_clean["Venta_Total"]) * 100).round(2)
        np.testing.assert_allclose(self.df_clean["Margen_Pct"], calculated_margin, rtol=1e-4)

    def test_audit_metrics_consistency(self):
        """Verifica la consistencia entre filas iniciales, descartadas y finales."""
        self.assertEqual(
            self.audit["filas_iniciales"],
            self.audit["filas_finales"] + self.audit["filas_descartadas_total"],
        )
        self.assertGreaterEqual(self.audit["tasa_retencion_pct"], 0.0)
        self.assertLessEqual(self.audit["tasa_retencion_pct"], 100.0)


class TestReportBuilder(unittest.TestCase):
    """Pruebas unitarias para la exportación de Excel y PDF en memoria."""

    def setUp(self):
        self.df_raw = generate_dirty_sales_data(filepath=None, n_rows=150, seed=77)
        clean_result = clean_sales_data(self.df_raw)
        self.df_clean = clean_result["cleaned_df"]
        self.audit = clean_result["audit_metrics"]
        self.kpis = clean_result["financial_kpis"]

    def test_export_to_excel_validity(self):
        """Verifica la generación de bytes válidos de Excel y sus hojas."""
        excel_bytes = export_to_excel_bytes(self.df_clean, self.audit, self.kpis)
        self.assertIsInstance(excel_bytes, bytes)
        self.assertGreater(len(excel_bytes), 5000, "El archivo Excel generado es demasiado pequeño.")

        # Cargar con openpyxl desde el buffer en memoria
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        sheet_names = wb.sheetnames
        self.assertIn("Resumen Ejecutivo", sheet_names)
        self.assertIn("Transacciones Limpias", sheet_names)
        self.assertIn("Analitica por Sucursal", sheet_names)

        # Verificar contenido en hoja limpia
        ws_data = wb["Transacciones Limpias"]
        self.assertGreaterEqual(ws_data.max_row, 10)
        self.assertGreaterEqual(ws_data.max_column, 8)

    def test_export_to_pdf_validity(self):
        """Verifica la generación de bytes de PDF con cabecera estándar y contenido binario."""
        pdf_bytes = export_to_pdf_bytes(self.audit, self.kpis)
        self.assertIsInstance(pdf_bytes, bytes)
        self.assertGreater(len(pdf_bytes), 10000, "El archivo PDF generado es demasiado pequeño.")

        # Verificar cabecera mágica de PDF (%PDF)
        self.assertTrue(pdf_bytes.startswith(b"%PDF"), "El archivo no cuenta con cabecera válida de PDF.")


if __name__ == "__main__":
    unittest.main(verbosity=2)
